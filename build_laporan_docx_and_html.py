# -*- coding: utf-8 -*-
"""
Script Penyusun Dokumen Laporan 3 dan 4 (SALUT)
Output:
  - Laporan_3_dan_4.docx (Microsoft Word Standar Akademik)
  - Laporan_3_dan_4.html (Web HTML Siap Cetak / Preview)
  - Laporan_3_dan_4.doc (Word-Compatible HTML)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = r"d:\proyek_applikasi\apps_sekolah"
IMG_DIR = os.path.join(BASE_DIR, "laporan_screenshots")
DOCX_PATH = os.path.join(BASE_DIR, "Laporan_3_dan_4.docx")
HTML_PATH = os.path.join(BASE_DIR, "Laporan_3_dan_4.html")
DOC_PATH = os.path.join(BASE_DIR, "Laporan_3_dan_4.doc")

def set_para_font(p, name='Times New Roman', size_pt=12, bold=False, italic=False, color_rgb=(0,0,0), space_after_pt=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.line_spacing = line_spacing
    for r in p.runs:
        r.font.name = name
        r.font.size = Pt(size_pt)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = RGBColor(*color_rgb)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    set_para_font(p, name='Times New Roman', size_pt=14, bold=True, color_rgb=(0, 51, 102), space_after_pt=12, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)
    return p

def add_subheading(doc, text):
    p = doc.add_paragraph()
    set_para_font(p, name='Times New Roman', size_pt=12, bold=True, color_rgb=(0, 0, 0), space_after_pt=4, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    return p

def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    set_para_font(p, name='Times New Roman', size_pt=12, bold=False, color_rgb=(0, 0, 0), space_after_pt=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p

def add_image_with_caption(doc, img_filename, caption_text, width_inches=6.0):
    img_path = os.path.join(IMG_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(14)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(10)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(80, 80, 80)
    else:
        print(f"Warning: Image {img_path} not found!")

def build_docx():
    print("Membuat dokumen Word (.docx)...")
    doc = docx.Document()
    
    # Page setup - Margins 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = False
        
        # Header & Footer
        header = section.header
        p_head = header.paragraphs[0]
        p_head.text = "Laporan Progres Sistem SALUT | Universitas Lampung"
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.runs[0].font.name = 'Times New Roman'
        p_head.runs[0].font.size = Pt(9)
        p_head.runs[0].font.italic = True
        p_head.runs[0].font.color.rgb = RGBColor(128, 128, 128)
        
    # ==========================================
    # LAPORAN 3
    # ==========================================
    add_heading_1(doc, "LAPORAN 3\nPENGEMBANGAN BACKEND API DAN KONFIGURASI DATABASE\nPADA SISTEM SALUT")
    
    # 1. Konfigurasi Lingkungan Backend
    add_subheading(doc, "1. Konfigurasi Lingkungan Backend (Laravel)")
    add_body_paragraph(doc, 
        "Pada tahap ini, dilakukan persiapan lingkungan backend menggunakan framework Laravel. "
        "Konfigurasi awal difokuskan pada pengaturan variabel lingkungan di dalam file .env, "
        "khususnya untuk menghubungkan aplikasi Laravel dengan database relasional (MySQL). "
        "Langkah ini krusial untuk memastikan aplikasi dapat berkomunikasi dengan sistem penyimpanan data, "
        "termasuk konfigurasi secret key untuk autentikasi JSON Web Token (JWT)."
    )
    add_image_with_caption(doc, "laporan3_01_env.png", "Gambar 3.1: Konfigurasi file .env untuk koneksi database MySQL dan JWT Auth", width_inches=6.0)
    
    # 2. Pembuatan Struktur Database
    add_subheading(doc, "2. Pembuatan Struktur Database (Migration dan Model)")
    add_body_paragraph(doc,
        "Setelah koneksi berhasil, tahap selanjutnya adalah merancang struktur database yang mendukung sistem "
        "multi-role (Admin, Guru, Siswa Aktif, Alumni). Pembuatan tabel dilakukan menggunakan fitur Migration bawaan "
        "Laravel agar skema database terstruktur, terdokumentasi, dan mudah di-reproduksi. "
        "Selain itu, dibuat juga file Model Eloquent (seperti User.php, AcademicRecord.php, Achievement.php, TracerStudy.php, "
        "dan GraduationDoc.php) untuk merepresentasikan tabel-tabel tersebut di dalam kode program."
    )
    add_image_with_caption(doc, "laporan3_02_migration_model.png", "Gambar 3.2: Struktur file Migration pada Laravel untuk tabel records dan relasi foreign key", width_inches=6.0)
    add_image_with_caption(doc, "laporan3_03_phpmyadmin.png", "Gambar 3.3: Tampilan struktur tabel database salut_db yang berhasil di-migrate pada phpMyAdmin", width_inches=6.0)
    
    # 3. Pembuatan Endpoint API dan Implementasi Autentikasi JWT
    add_subheading(doc, "3. Pembuatan Endpoint API dan Implementasi Autentikasi JWT")
    add_body_paragraph(doc,
        "Sistem SALUT menggunakan arsitektur Client-Server, sehingga diperlukan pembuatan API (Application Programming Interface). "
        "Untuk menjaga keamanan data dan komunikasi, diimplementasikan autentikasi berbasis token menggunakan JWT (JSON Web Token). "
        "Endpoint login publik disediakan pada rute POST /api/login, sedangkan rute modul profil, prestasi siswa, arsip digital kelulusan, "
        "dan tracer study diproteksi secara ketat menggunakan middleware jwt.auth yang mewajibkan penyertaan Bearer Token."
    )
    add_image_with_caption(doc, "laporan3_04_routes_api.png", "Gambar 3.4: Tampilan routing RESTful API dan proteksi middleware jwt.auth pada file routes/api.php", width_inches=6.0)
    
    # 4. Pengujian API (API Testing)
    add_subheading(doc, "4. Pengujian API (API Testing)")
    add_body_paragraph(doc,
        "Untuk memastikan endpoint API dan autentikasi JWT berjalan dengan baik sebelum diintegrasikan ke aplikasi mobile, "
        "dilakukan pengujian menggunakan perangkat lunak API Client (Postman / Thunder Client). "
        "Pengujian dilakukan dengan mengirimkan request login berupa kredensial NISN dan password, "
        "dan server berhasil merespons dengan status 200 OK serta mengembalikan payload token akses JWT beserta objek data profil pengguna."
    )
    add_image_with_caption(doc, "laporan3_05_api_testing.png", "Gambar 3.5: Hasil pengujian API login dengan respons status 200 OK dan token JWT yang valid", width_inches=6.0)
    
    # 5. Penyimpanan Repositori Backend ke GitHub
    add_subheading(doc, "5. Penyimpanan Repositori Backend ke GitHub")
    add_body_paragraph(doc,
        "Setelah kerangka dasar backend selesai dan diuji secara komprehensif, seluruh perubahan kode (source code) disimpan "
        "ke version control system GitHub pada repositori WEMVIAFRIALDO/apps_Sekolah-. "
        "Langkah ini bertujuan untuk mengamankan data proyek, melacak riwayat perubahan kode (commit history), "
        "dan memudahkan kolaborasi pengembangan di masa mendatang."
    )
    add_image_with_caption(doc, "laporan3_06_git_push_backend.png", "Gambar 3.6: Eksekusi perintah git commit dan git push pada terminal untuk backend Laravel", width_inches=6.0)
    add_image_with_caption(doc, "laporan3_07_github_backend.png", "Gambar 3.7: Tampilan repositori Backend SALUT yang telah tersinkronisasi di website GitHub", width_inches=6.0)
    
    # Page Break antara Laporan 3 dan Laporan 4
    doc.add_page_break()
    
    # ==========================================
    # LAPORAN 4
    # ==========================================
    add_heading_1(doc, "LAPORAN 4\nPENGEMBANGAN ANTARMUKA MOBILE (FLUTTER) DAN INTEGRASI API")
    
    # 1. Pembuatan Antarmuka Pengguna (UI) Utama
    add_subheading(doc, "1. Pembuatan Antarmuka Pengguna (UI) Utama")
    add_body_paragraph(doc,
        "Pada tahap ini, pengembangan beralih ke sisi frontend mobile menggunakan framework Flutter. "
        "Langkah pertama adalah menerjemahkan desain (mockup) menjadi kode Dart untuk membentuk antarmuka pengguna (UI). "
        "Fokus utama difokuskan pada perancangan halaman autentikasi (Login Page) dengan form validasi input NISN dan kata sandi, "
        "serta halaman utama (Dashboard Page) yang responsif dan dirancang dengan palet warna brand modern (#1E3A5F)."
    )
    add_image_with_caption(doc, "laporan4_01_login_page_code.png", "Gambar 4.1: Potongan kode antarmuka halaman login pada file login_page.dart", width_inches=6.0)
    
    # 2. Integrasi HTTP Request dan Pengelolaan Token JWT
    add_subheading(doc, "2. Integrasi HTTP Request dan Pengelolaan Token JWT")
    add_body_paragraph(doc,
        "Agar aplikasi mobile tidak sekadar menampilkan antarmuka statis, dilakukan integrasi dengan Backend API "
        "yang telah dibuat pada Laporan 3. Aplikasi Flutter dikonfigurasi menggunakan pustaka HTTP client Dio "
        "dengan arsitektur Interceptor pada ApiService. Saat pengguna berhasil login, Bearer Token JWT disimpan "
        "ke penyimpanan lokal perangkat secara aman (secure storage) dan disisipkan secara otomatis pada setiap request ke endpoint terproteksi."
    )
    add_image_with_caption(doc, "laporan4_02_api_service_code.png", "Gambar 4.2: Implementasi HTTP Client Dio dan Interceptor token JWT pada api_service.dart", width_inches=6.0)
    
    # 3. Simulasi dan Pengujian Aplikasi pada Perangkat
    add_subheading(doc, "3. Simulasi dan Pengujian Aplikasi pada Perangkat")
    add_body_paragraph(doc,
        "Setelah integrasi selesai, aplikasi dijalankan (Run/Debug) untuk simulasi operasional. "
        "Pengujian dilakukan dengan memasukkan kredensial login akun pengujian untuk memverifikasi apakah aplikasi berhasil "
        "terhubung ke server lokal dan berpindah dari halaman Login menuju halaman Dashboard. "
        "Pengujian berhasil mendemonstrasikan dua skenario role: Siswa Aktif (Budi Santoso) dan Alumni (Siti Rahayu)."
    )
    add_image_with_caption(doc, "laporan4_03_login_mobile_app.png", "Gambar 4.3: Tampilan antarmuka Login aplikasi SALUT saat berjalan di perangkat", width_inches=3.4)
    add_image_with_caption(doc, "laporan4_04_dashboard_siswa_app.png", "Gambar 4.4: Tampilan antarmuka Dashboard untuk peran Siswa Aktif (fitur Prestasi dan Rapor)", width_inches=3.4)
    add_image_with_caption(doc, "laporan4_05_dashboard_alumni_app.png", "Gambar 4.5: Tampilan antarmuka Dashboard untuk peran Alumni (fitur Arsip Digital dan Tracer Study)", width_inches=3.4)
    
    # 4. Implementasi Logika Dynamic UI Rendering (Role-Based)
    add_subheading(doc, "4. Implementasi Logika Dynamic UI Rendering (Role-Based)")
    add_body_paragraph(doc,
        "Sesuai dengan aturan bisnis sistem SALUT, tampilan antarmuka dirancang dinamis menyesuaikan role (peran) pengguna "
        "yang didapatkan dari payload autentikasi JWT. Diterapkan logika kondisional berbasis state management GetX: "
        "jika pengguna terdeteksi sebagai Siswa Aktif, maka navigasi dan menu yang muncul difokuskan pada manajemen Prestasi dan Nilai Rapor; "
        "sedangkan jika pengguna terdeteksi sebagai Alumni, menu secara otomatis berganti menjadi Arsip Digital Ijazah/SKL dan Kuesioner Tracer Study."
    )
    add_image_with_caption(doc, "laporan4_06_role_logic_code.png", "Gambar 4.6: Potongan kode logika pengkondisian tampilan menu per-role pada dashboard_page.dart", width_inches=6.0)
    
    # 5. Penyimpanan Repositori Frontend ke GitHub
    add_subheading(doc, "5. Penyimpanan Repositori Frontend ke GitHub")
    add_body_paragraph(doc,
        "Sebagai langkah penutup tahap pengembangan frontend mobile, pembaruan kode aplikasi Flutter disimpan "
        "dan disinkronisasikan ke repositori GitHub WEMVIAFRIALDO/apps_Sekolah-. "
        "Dengan tersimpannya seluruh kode sumber ini, repositori monorepo proyek SALUT telah mencakup backend Laravel, "
        "web admin panel, serta aplikasi mobile Flutter secara terpadu."
    )
    add_image_with_caption(doc, "laporan4_07_git_push_frontend.png", "Gambar 4.7: Tampilan terminal saat melakukan commit dan push kode frontend Flutter ke GitHub", width_inches=6.0)
    add_image_with_caption(doc, "laporan4_08_github_frontend.png", "Gambar 4.8: Tampilan repositori Frontend SALUT yang berhasil di-push ke GitHub", width_inches=6.0)
    
    doc.save(DOCX_PATH)
    print("Dokumen Word berhasil disimpan di:", DOCX_PATH)

def build_html_and_doc():
    print("Membuat file HTML dan DOC...")
    html_content = """<!DOCTYPE html>
<html lang="id">
<head>
<meta charset="utf-8">
<title>Laporan Progres 3 dan 4 - Sistem SALUT</title>
<style>
  @page {
    size: A4;
    margin: 2.5cm;
  }
  body { 
    font-family: 'Times New Roman', serif; 
    font-size: 12pt; 
    line-height: 1.6; 
    padding: 40px; 
    max-width: 860px; 
    margin: auto; 
    color: #111;
    background-color: #fff;
  }
  h1 { 
    font-size: 15pt; 
    text-align: center; 
    font-weight: bold; 
    margin-top: 35px; 
    margin-bottom: 25px;
    color: #003366;
    line-height: 1.3;
  }
  h2 {
    font-size: 12pt;
    font-weight: bold;
    margin-top: 25px;
    margin-bottom: 8px;
    color: #000;
  }
  p { 
    text-align: justify; 
    margin-bottom: 12px;
    text-indent: 35px;
  }
  figure {
    text-align: center;
    margin: 20px 0;
  }
  figure img {
    max-width: 95%;
    height: auto;
    border: 1px solid #ddd;
    border-radius: 4px;
    box-shadow: 0 2px 6px rgba(0,0,0,0.1);
  }
  figure.mobile-frame img {
    max-width: 380px;
  }
  figcaption {
    font-size: 10pt;
    font-style: italic;
    color: #555;
    margin-top: 8px;
    text-align: center;
  }
  .page-break {
    page-break-before: always;
    margin-top: 50px;
    border-top: 2px dashed #ccc;
    padding-top: 40px;
  }
</style>
</head>
<body>

<h1>LAPORAN 3<br>PENGEMBANGAN BACKEND API DAN KONFIGURASI DATABASE PADA SISTEM SALUT</h1>

<h2>1. Konfigurasi Lingkungan Backend (Laravel)</h2>
<p>Pada tahap ini, dilakukan persiapan lingkungan backend menggunakan framework Laravel. Konfigurasi awal difokuskan pada pengaturan variabel lingkungan di dalam file .env, khususnya untuk menghubungkan aplikasi Laravel dengan database relasional (MySQL). Langkah ini krusial untuk memastikan aplikasi dapat berkomunikasi dengan sistem penyimpanan data, termasuk konfigurasi secret key untuk autentikasi JSON Web Token (JWT).</p>
<figure>
  <img src="laporan_screenshots/laporan3_01_env.png" alt="Konfigurasi .env Backend">
  <figcaption>Gambar 3.1: Konfigurasi file .env untuk koneksi database MySQL dan JWT Auth</figcaption>
</figure>

<h2>2. Pembuatan Struktur Database (Migration dan Model)</h2>
<p>Setelah koneksi berhasil, tahap selanjutnya adalah merancang struktur database yang mendukung sistem multi-role (Admin, Guru, Siswa Aktif, Alumni). Pembuatan tabel dilakukan menggunakan fitur Migration bawaan Laravel agar skema database lebih terstruktur, aman, dan mudah di-reproduksi. Selain itu, dibuat juga file Model Eloquent (seperti User.php, AcademicRecord.php, Achievement.php, TracerStudy.php, dan GraduationDoc.php) untuk merepresentasikan tabel-tabel tersebut di dalam kode.</p>
<figure>
  <img src="laporan_screenshots/laporan3_02_migration_model.png" alt="Kode Migration dan Model Laravel">
  <figcaption>Gambar 3.2: Struktur file Migration pada Laravel untuk tabel records dan relasi foreign key</figcaption>
</figure>
<figure>
  <img src="laporan_screenshots/laporan3_03_phpmyadmin.png" alt="Tampilan Tabel Database di phpMyAdmin">
  <figcaption>Gambar 3.3: Tampilan struktur tabel database salut_db yang berhasil di-migrate pada phpMyAdmin</figcaption>
</figure>

<h2>3. Pembuatan Endpoint API dan Implementasi Autentikasi JWT</h2>
<p>Sistem SALUT menggunakan arsitektur Client-Server, sehingga diperlukan pembuatan API (Application Programming Interface). Untuk menjaga keamanan data dan komunikasi, diimplementasikan autentikasi berbasis token menggunakan JWT (JSON Web Token). Endpoint login publik disediakan pada rute POST /api/login, sedangkan rute modul profil, prestasi siswa, arsip digital kelulusan, dan tracer study diproteksi menggunakan middleware jwt.auth yang mewajibkan penyertaan Bearer Token.</p>
<figure>
  <img src="laporan_screenshots/laporan3_04_routes_api.png" alt="Kode Routing API Laravel">
  <figcaption>Gambar 3.4: Tampilan routing RESTful API dan proteksi middleware jwt.auth pada file routes/api.php</figcaption>
</figure>

<h2>4. Pengujian API (API Testing)</h2>
<p>Untuk memastikan endpoint API dan autentikasi JWT berjalan dengan baik sebelum diintegrasikan ke aplikasi mobile, dilakukan pengujian menggunakan perangkat lunak API Client (seperti Postman atau Thunder Client). Pengujian dilakukan dengan mengirimkan request login berupa kredensial NISN dan password, dan server berhasil merespons dengan mengembalikan status 200 OK serta token akses JWT yang valid.</p>
<figure>
  <img src="laporan_screenshots/laporan3_05_api_testing.png" alt="Pengujian API Client 200 OK">
  <figcaption>Gambar 3.5: Hasil pengujian API login dengan respons status 200 OK dan token JWT yang valid</figcaption>
</figure>

<h2>5. Penyimpanan Repositori Backend ke GitHub</h2>
<p>Setelah kerangka dasar backend selesai dan diuji, seluruh perubahan kode (source code) disimpan ke version control system GitHub pada repositori WEMVIAFRIALDO/apps_Sekolah-. Hal ini bertujuan untuk mengamankan data proyek, mendokumentasikan riwayat perubahan kode, dan memudahkan proses integrasi antar modul.</p>
<figure>
  <img src="laporan_screenshots/laporan3_06_git_push_backend.png" alt="Terminal Git Push Backend">
  <figcaption>Gambar 3.6: Eksekusi perintah git commit dan git push pada terminal untuk backend Laravel</figcaption>
</figure>
<figure>
  <img src="laporan_screenshots/laporan3_07_github_backend.png" alt="Repositori Backend di GitHub">
  <figcaption>Gambar 3.7: Tampilan repositori Backend SALUT yang telah tersinkronisasi di website GitHub</figcaption>
</figure>

<div class="page-break"></div>

<h1>LAPORAN 4<br>PENGEMBANGAN ANTARMUKA MOBILE (FLUTTER) DAN INTEGRASI API</h1>

<h2>1. Pembuatan Antarmuka Pengguna (UI) Utama</h2>
<p>Pada tahap ini, pengembangan beralih ke sisi frontend mobile menggunakan Flutter. Langkah pertama adalah menerjemahkan desain (mockup) menjadi kode Dart untuk membentuk antarmuka pengguna (UI). Fokus utama difokuskan pada perancangan halaman autentikasi (Login Page) dengan form validasi NISN dan kata sandi, serta halaman utama (Dashboard Page) yang responsif dan dirancang dengan palet warna brand modern (#1E3A5F).</p>
<figure>
  <img src="laporan_screenshots/laporan4_01_login_page_code.png" alt="Kode Flutter Login Page">
  <figcaption>Gambar 4.1: Potongan kode antarmuka halaman login pada file login_page.dart</figcaption>
</figure>

<h2>2. Integrasi HTTP Request dan Pengelolaan Token JWT</h2>
<p>Agar aplikasi mobile tidak sekadar menampilkan UI statis, dilakukan integrasi dengan Backend API yang telah dibuat pada Laporan 3. Aplikasi Flutter dikonfigurasi menggunakan pustaka HTTP client Dio dengan arsitektur Interceptor pada ApiService. Saat pengguna berhasil login, token JWT disimpan ke penyimpanan lokal perangkat secara aman (secure storage) dan disisipkan secara otomatis pada setiap request ke server.</p>
<figure>
  <img src="laporan_screenshots/laporan4_02_api_service_code.png" alt="Kode ApiService Flutter Dio JWT">
  <figcaption>Gambar 4.2: Implementasi HTTP Client Dio dan Interceptor token JWT pada api_service.dart</figcaption>
</figure>

<h2>3. Simulasi dan Pengujian Aplikasi pada Perangkat</h2>
<p>Setelah integrasi selesai, aplikasi dijalankan (Run/Debug) untuk simulasi operasional. Pengujian dilakukan dengan memasukkan kredensial login untuk memverifikasi apakah aplikasi berhasil terhubung ke server lokal dan berpindah dari halaman Login menuju Dashboard. Pengujian berhasil mendemonstrasikan dua skenario role: Siswa Aktif (Budi Santoso) dan Alumni (Siti Rahayu).</p>
<figure class="mobile-frame">
  <img src="laporan_screenshots/laporan4_03_login_mobile_app.png" alt="Aplikasi Mobile Halaman Login">
  <figcaption>Gambar 4.3: Tampilan antarmuka Login aplikasi SALUT saat berjalan di perangkat</figcaption>
</figure>
<figure class="mobile-frame">
  <img src="laporan_screenshots/laporan4_04_dashboard_siswa_app.png" alt="Dashboard Mobile Siswa Aktif">
  <figcaption>Gambar 4.4: Tampilan antarmuka Dashboard untuk peran Siswa Aktif (fitur Prestasi dan Rapor)</figcaption>
</figure>
<figure class="mobile-frame">
  <img src="laporan_screenshots/laporan4_05_dashboard_alumni_app.png" alt="Dashboard Mobile Alumni">
  <figcaption>Gambar 4.5: Tampilan antarmuka Dashboard untuk peran Alumni (fitur Arsip Digital dan Tracer Study)</figcaption>
</figure>

<h2>4. Implementasi Logika Dynamic UI Rendering (Role-Based)</h2>
<p>Sesuai dengan aturan bisnis sistem SALUT, tampilan antarmuka dirancang dinamis menyesuaikan role (peran) pengguna yang didapatkan dari payload autentikasi JWT. Diterapkan logika kondisional berbasis state management GetX: jika pengguna terdeteksi sebagai Siswa Aktif, maka navigasi dan menu difokuskan pada Prestasi dan Rapor; sedangkan jika terdeteksi sebagai Alumni, menu otomatis berganti menjadi Arsip Digital Ijazah/SKL dan Kuesioner Tracer Study.</p>
<figure>
  <img src="laporan_screenshots/laporan4_06_role_logic_code.png" alt="Kode Dynamic UI Role Dashboard">
  <figcaption>Gambar 4.6: Potongan kode logika pengkondisian tampilan menu per-role pada dashboard_page.dart</figcaption>
</figure>

<h2>5. Penyimpanan Repositori Frontend ke GitHub</h2>
<p>Sebagai langkah penutup tahap pengembangan frontend mobile, pembaruan kode aplikasi Flutter disimpan dan disinkronisasikan ke repositori GitHub WEMVIAFRIALDO/apps_Sekolah-. Dengan tersimpannya seluruh kode sumber ini, repositori monorepo proyek SALUT telah mencakup backend Laravel, web admin panel, serta aplikasi mobile Flutter secara terpadu.</p>
<figure>
  <img src="laporan_screenshots/laporan4_07_git_push_frontend.png" alt="Terminal Git Push Frontend Flutter">
  <figcaption>Gambar 4.7: Tampilan terminal saat melakukan commit dan push kode frontend Flutter ke GitHub</figcaption>
</figure>
<figure>
  <img src="laporan_screenshots/laporan4_08_github_frontend.png" alt="Repositori Frontend di GitHub">
  <figcaption>Gambar 4.8: Tampilan repositori Frontend SALUT yang berhasil di-push ke GitHub</figcaption>
</figure>

</body>
</html>
"""
    with open(HTML_PATH, "w", encoding="utf-8") as f:
        f.write(html_content)
    print("File HTML berhasil disimpan di:", HTML_PATH)

    with open(DOC_PATH, "w", encoding="utf-8") as f:
        f.write(html_content)
    print("File DOC (Word Compatible) berhasil disimpan di:", DOC_PATH)

if __name__ == "__main__":
    build_docx()
    build_html_and_doc()
    print("Semua file dokumen laporan berhasil disusun lengkap!")
